"""Monta a planilha Excel (.xlsx) de um participante.

Regra fixa do projeto: os dados nunca ficam "soltos" num único log — cada
participante tem seu próprio arquivo, com abas separadas por tipo de dado, e
dentro delas sempre desagregados por fase (e por bin de 10s), nunca só o
total da sessão. Ver `experiment/analytics.py` para o cálculo das agregações.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = "2563EB"


def _write_header(ws, headers: list[str]) -> None:
    from openpyxl.styles import PatternFill

    ws.append(headers)
    header_row = ws.max_row
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = f"A{header_row + 1}"


def _autosize(ws, min_width: int = 10, max_width: int = 40) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        longest = max(
            (len(str(row[col_idx - 1].value)) for row in ws.iter_rows() if row[col_idx - 1].value is not None),
            default=0,
        )
        ws.column_dimensions[letter].width = max(min_width, min(longest + 2, max_width))


def _sheet_resumo(wb: Workbook, participant, sessions) -> None:
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Participante"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["ID (interno)", str(participant.id)])
    ws.append(["Identificador (operador)", participant.external_id])
    ws.append(["Idade", participant.age if participant.age is not None else ""])
    ws.append(["Sexo", participant.get_sex_display() if participant.sex else ""])
    ws.append(["Cadastrado em", participant.created_at.strftime("%d/%m/%Y %H:%M")])
    ws.append([])
    sessoes_title_row = ws.max_row + 1
    ws.append(["Sessões"])
    ws.cell(row=sessoes_title_row, column=1).font = Font(bold=True, size=12)
    _write_header(
        ws,
        ["session_id", "grupo", "modo_dev", "status", "criada_em", "iniciada_em", "concluida_em"],
    )
    for s in sessions:
        ws.append(
            [
                str(s.id),
                s.group,
                "sim" if s.dev_mode else "não",
                s.status,
                s.created_at.strftime("%d/%m/%Y %H:%M") if s.created_at else "",
                s.started_at.strftime("%d/%m/%Y %H:%M") if s.started_at else "",
                s.finished_at.strftime("%d/%m/%Y %H:%M") if s.finished_at else "",
            ]
        )
    _autosize(ws)


def _sheet_eventos(wb: Workbook, sessions) -> None:
    ws = wb.create_sheet("Eventos")
    _write_header(
        ws,
        [
            "session_id",
            "grupo",
            "fase",
            "tipo_evento",
            "botao",
            "t_ms",
            "pontos_delta",
            "pontos_total",
            "som_ms",
        ],
    )
    for s in sessions:
        for e in s.events.all():
            ws.append(
                [
                    str(s.id),
                    s.group,
                    e.phase,
                    e.event_type,
                    e.button_role,
                    e.t_ms,
                    e.points_delta,
                    e.points_total,
                    e.sound_ms if e.sound_ms is not None else "",
                ]
            )
    _autosize(ws)


def _sheet_por_fase(wb: Workbook, sessions) -> None:
    ws = wb.create_sheet("Por Fase")
    _write_header(
        ws,
        [
            "session_id",
            "grupo",
            "fase",
            "pontos_inicio_fase",
            "pontos_fim_fase",
            "pontos_delta_fase",
            "duracao_s",
            "botao",
            "respostas",
            "respostas_por_min",
            "reforcos",
            "reforcos_por_min",
            "custos",
            "pontos_delta_botao",
        ],
    )
    for s in sessions:
        phase_stats = {p.phase: p for p in s.phase_stats.all()}
        for bp in s.button_phase_stats.all():
            ps = phase_stats.get(bp.phase)
            ws.append(
                [
                    str(s.id),
                    s.group,
                    bp.phase,
                    ps.points_start if ps else "",
                    ps.points_end if ps else "",
                    ps.points_delta if ps else "",
                    ps.duration_s if ps else "",
                    bp.button_role,
                    bp.response_count,
                    round(bp.responses_per_minute, 2),
                    bp.reinforcement_count,
                    round(bp.reinforcements_per_minute, 2),
                    bp.cost_count,
                    bp.points_delta,
                ]
            )
    _autosize(ws)


def _sheet_bins(wb: Workbook, sessions) -> None:
    ws = wb.create_sheet("Bins 10s")
    _write_header(
        ws,
        [
            "session_id",
            "grupo",
            "fase",
            "botao",
            "bin_index",
            "bin_inicio_s",
            "bin_fim_s",
            "respostas",
            "reforcos",
            "custos",
        ],
    )
    for s in sessions:
        for b in s.phase_bins.all():
            ws.append(
                [
                    str(s.id),
                    s.group,
                    b.phase,
                    b.button_role,
                    b.bin_index,
                    b.bin_start_s,
                    b.bin_end_s,
                    b.response_count,
                    b.reinforcement_count,
                    b.cost_count,
                ]
            )
    _autosize(ws)


def participant_filename(participant) -> str:
    """Nome de arquivo único mesmo se dois participantes tiverem o mesmo
    external_id (o operador pode digitar o mesmo identificador por engano —
    já aconteceu nos dados reais deste projeto)."""
    short_id = str(participant.id).split("-")[0]
    safe_external_id = "".join(
        c if c.isalnum() or c in "-_" else "_" for c in participant.external_id
    ) or "participante"
    return f"participante_{safe_external_id}_{short_id}.xlsx"


def build_participant_workbook(participant) -> bytes:
    """Monta o .xlsx completo de um participante: 1 aba por tipo de dado,
    sempre desagregado por fase/bin dentro de cada aba (nunca só o total)."""
    sessions = list(
        participant.sessions.all()
        .prefetch_related("events", "phase_stats", "button_phase_stats", "phase_bins")
        .order_by("created_at")
    )
    wb = Workbook()
    _sheet_resumo(wb, participant, sessions)
    _sheet_eventos(wb, sessions)
    _sheet_por_fase(wb, sessions)
    _sheet_bins(wb, sessions)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
